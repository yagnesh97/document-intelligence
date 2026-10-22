from io import BytesIO

from openpyxl import Workbook

from app.services.parsers.excel import ExcelParser


def create_excel_file() -> bytes:
    workbook = Workbook()

    worksheet = workbook.active
    worksheet.title = "Transactions"

    worksheet.append(["Date", "Description", "Amount"])
    worksheet.append(["2026-01-01", "Salary", 5000])
    worksheet.append(["2026-01-02", "Coffee", -5])

    second_sheet = workbook.create_sheet("Summary")
    second_sheet.append(["Total", 4995])

    output = BytesIO()
    workbook.save(output)

    return output.getvalue()


def test_excel_parser_extracts_worksheets_and_content():
    content = create_excel_file()

    parser = ExcelParser()

    result = parser.parse(
        content=content,
        filename="transactions.xlsx",
        content_type=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
    )

    assert "Transactions" in result.content
    assert "Date" in result.content
    assert "Salary" in result.content
    assert "5000" in result.content

    assert "Summary" in result.content
    assert "4995" in result.content

    assert result.metadata["filename"] == "transactions.xlsx"
    assert result.metadata["worksheet_count"] == 2
    assert result.metadata["worksheets"] == [
        "Transactions",
        "Summary",
    ]


def test_excel_parser_ignores_empty_rows():
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Data"

    worksheet.append(["Name", "Amount"])
    worksheet.append([None, None])
    worksheet.append(["Coffee", 5])

    output = BytesIO()
    workbook.save(output)

    parser = ExcelParser()

    result = parser.parse(
        content=output.getvalue(),
        filename="test.xlsx",
    )

    assert "Name\tAmount" in result.content
    assert "Coffee\t5" in result.content

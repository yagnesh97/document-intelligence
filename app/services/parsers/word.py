from io import BytesIO

from docx import Document as WordDocument

from app.services.parsers.base import BaseParser, ParsedDocument


class WordParser(BaseParser):
    """Extract text and tables from DOCX documents."""

    def parse(
        self,
        content: bytes,
        filename: str,
        content_type: str | None = None,
    ) -> ParsedDocument:
        document = WordDocument(BytesIO(content))
        sections: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                sections.append(text)

        table_count = 0

        for table_index, table in enumerate(document.tables, start=1):
            rows = [
                "\t".join(cell.text.strip() for cell in row.cells) for row in table.rows
            ]
            if rows:
                sections.append(f"[Table {table_index}]\n" + "\n".join(rows))
                table_count += 1

        extracted_text = "\n\n".join(sections)

        if not extracted_text.strip():
            raise ValueError(f"No text could be extracted from '{filename}'")

        return ParsedDocument(
            content=extracted_text,
            metadata={
                "filename": filename,
                "content_type": content_type,
                "table_count": table_count,
            },
        )

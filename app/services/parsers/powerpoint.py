from io import BytesIO

from pptx import Presentation

from app.services.parsers.base import BaseParser, ParsedDocument


class PowerPointParser(BaseParser):
    """Extract text from PowerPoint slides and tables."""

    def parse(
        self,
        content: bytes,
        filename: str,
        content_type: str | None = None,
    ) -> ParsedDocument:
        presentation = Presentation(BytesIO(content))
        slides_text: list[str] = []

        for slide_number, slide in enumerate(presentation.slides, start=1):
            elements: list[str] = []

            for shape in slide.shapes:
                if shape.has_text_frame:
                    text = shape.text.strip()
                    if text:
                        elements.append(text)

                if shape.has_table:
                    rows = [
                        "\t".join(cell.text.strip() for cell in row.cells)
                        for row in shape.table.rows
                    ]
                    if rows:
                        elements.append("\n".join(rows))

            if elements:
                slides_text.append(f"[Slide {slide_number}]\n" + "\n".join(elements))

        extracted_text = "\n\n".join(slides_text)

        if not extracted_text.strip():
            raise ValueError(f"No text could be extracted from '{filename}'")

        return ParsedDocument(
            content=extracted_text,
            metadata={
                "filename": filename,
                "content_type": content_type,
                "slide_count": len(presentation.slides),
            },
        )

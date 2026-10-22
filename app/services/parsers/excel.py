from io import BytesIO
from pathlib import Path

import xlrd
from openpyxl import load_workbook

from app.services.parsers.base import BaseParser, ParsedDocument


class ExcelParser(BaseParser):
    """Extract worksheet data from XLS and XLSX files."""

    def parse(
        self,
        content: bytes,
        filename: str,
        content_type: str | None = None,
    ) -> ParsedDocument:
        extension = Path(filename).suffix.lower()
        sections: list[str] = []
        worksheets: list[str] = []

        if extension == ".xls":
            workbook = xlrd.open_workbook(file_contents=content)

            try:
                for sheet in workbook.sheets():
                    worksheets.append(sheet.name)
                    rows: list[str] = []

                    for row_index in range(sheet.nrows):
                        values = [
                            str(value).strip() for value in sheet.row_values(row_index)
                        ]
                        if any(values):
                            rows.append("\t".join(values))

                    if rows:
                        sections.append(f"[Sheet: {sheet.name}]\n" + "\n".join(rows))
            finally:
                workbook.release_resources()

        elif extension == ".xlsx":
            workbook = load_workbook(
                BytesIO(content),
                read_only=True,
                data_only=True,
            )

            try:
                for worksheet in workbook.worksheets:
                    worksheets.append(worksheet.title)
                    rows: list[str] = []

                    for row in worksheet.iter_rows(values_only=True):
                        values = ["" if value is None else str(value) for value in row]
                        if any(value.strip() for value in values):
                            rows.append("\t".join(values))

                    if rows:
                        sections.append(
                            f"[Sheet: {worksheet.title}]\n" + "\n".join(rows)
                        )
            finally:
                workbook.close()

        else:
            raise ValueError(f"Unsupported Excel format: {extension}")

        extracted_text = "\n\n".join(sections)

        if not extracted_text.strip():
            raise ValueError(f"No text could be extracted from '{filename}'")

        return ParsedDocument(
            content=extracted_text,
            metadata={
                "filename": filename,
                "content_type": content_type,
                "worksheet_count": len(worksheets),
                "worksheets": worksheets,
            },
        )
